"""Timeline ground-truth import and persistence."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

from bas_har.schema.activity_schema import ActivityId, TimelineRecord
from bas_har.studio.registry import ActivityRegistry
from bas_har.studio.takes import list_takes

_TIME_PATTERN = re.compile(r"^(?P<hours>\d+):(?P<minutes>[0-5]\d):(?P<seconds>[0-5]\d(?:\.\d+)?)$")
REQUIRED_COLUMNS = {
    "id",
    "take_id",
    "start_s",
    "end_s",
    "expected_step_id",
    "observed_action",
    "result",
}


def parse_time(value: Any) -> float:
    if isinstance(value, bool):
        raise ValueError("timeline time must be a number or HH:MM:SS.mmm")
    if isinstance(value, (int, float)):
        result = float(value)
        if result < 0:
            raise ValueError("timeline time cannot be negative")
        return result
    if not isinstance(value, str):
        raise ValueError("timeline time must be a number or HH:MM:SS.mmm")
    stripped = value.strip()
    if not stripped:
        raise ValueError("timeline time cannot be empty")
    try:
        result = float(stripped)
    except ValueError:
        match = _TIME_PATTERN.fullmatch(stripped)
        if match is None:
            raise ValueError(f"invalid timeline time: {value!r}") from None
        result = (
            int(match.group("hours")) * 3600
            + int(match.group("minutes")) * 60
            + float(match.group("seconds"))
        )
    if result < 0:
        raise ValueError("timeline time cannot be negative")
    return result


def read_timeline(path: Path) -> list[TimelineRecord]:
    if not path.is_file():
        raise FileNotFoundError(f"timeline file not found: {path}")
    if path.suffix.lower() == ".csv":
        rows = _read_csv(path)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(path)
    else:
        raise ValueError(f"unsupported timeline type: {path.suffix or 'none'}")
    if not rows:
        raise ValueError("timeline file contains no records")
    columns = set(rows[0])
    missing = REQUIRED_COLUMNS - columns
    if missing:
        raise ValueError(f"timeline is missing required columns: {sorted(missing)}")
    records: list[TimelineRecord] = []
    for row in rows:
        values = dict(row)
        values["start_s"] = parse_time(values.get("start_s"))
        values["end_s"] = parse_time(values.get("end_s"))
        values["object_ids"] = _split_ids(values.get("object_ids", ""))
        values["region_ids"] = _split_ids(values.get("region_ids", ""))
        records.append(TimelineRecord.model_validate(values))
    return records


def import_timeline(
    registry: ActivityRegistry,
    activity_id: ActivityId | str,
    source: Path,
) -> list[TimelineRecord]:
    records = read_timeline(source)
    takes = {take.take_id: take for take in list_takes(registry, activity_id)}
    for record in records:
        take = takes.get(record.take_id)
        if take is None:
            raise ValueError(f"timeline references unknown take: {record.take_id}")
        if take.duration_s is not None and record.end_s > take.duration_s + 0.05:
            raise ValueError(
                f"timeline record {record.record_id} ends after take {record.take_id} duration"
            )
    manifest = registry.load(activity_id)
    timeline_path = registry.package_dir(manifest.activity_id) / "timeline.jsonl"
    timeline_path.write_text(
        "".join(record.model_dump_json(by_alias=True) + "\n" for record in records),
        encoding="utf-8",
    )
    return records


def list_timeline(
    registry: ActivityRegistry,
    activity_id: ActivityId | str,
) -> list[TimelineRecord]:
    manifest = registry.load(activity_id)
    timeline_path = registry.package_dir(manifest.activity_id) / "timeline.jsonl"
    if not timeline_path.is_file():
        return []
    records: list[TimelineRecord] = []
    for line in timeline_path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            records.append(TimelineRecord.model_validate(json.loads(line)))
    return records


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "Excel import requires the optional studio dependency: pip install -e .[studio]"
        ) from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Timeline" not in workbook.sheetnames:
        raise ValueError("Excel workbook must contain a Timeline sheet")
    sheet = workbook["Timeline"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
    return [
        dict(zip(headers, row, strict=False))
        for row in rows
        if any(value is not None for value in row)
    ]


def _split_ids(value: Any) -> list[str]:
    if value in (None, ""):
        return []
    if not isinstance(value, str):
        value = str(value)
    return [item.strip() for item in value.split(",") if item.strip()]


__all__ = ["import_timeline", "list_timeline", "parse_time", "read_timeline"]
